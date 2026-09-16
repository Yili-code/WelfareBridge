"""共用解析工具：HTML → 純文字、標籤欄位抽取、PDF / DOCX 文字、附件連結。

只做「原文 → 文字」的轉換，不做任何猜測或補值。
"""

from __future__ import annotations

import io
import re
from urllib.parse import urljoin

from bs4 import BeautifulSoup, Tag

BLOCK_TAGS = {"p", "div", "li", "tr", "br", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section", "article", "dd", "dt"}
# 注意：不能丟 <form>——ASP.NET 網站（例如內政部住宅補貼平台）整個內容區都包在 <form> 裡
DROP_TAGS = {"script", "style", "noscript", "iframe", "svg", "nav", "header", "footer"}
ATTACHMENT_EXT = re.compile(r"\.(pdf|docx?|odt|ods|xlsx?|pptx?|zip|rar|7z)(\?|$)", re.I)


def make_soup(html: str) -> BeautifulSoup:
    # 環境沒有 lxml wheel，統一使用內建 html.parser
    return BeautifulSoup(html, "html.parser")


def page_title(html: str) -> str:
    soup = make_soup(html)
    if soup.title and soup.title.string:
        return clean_text(soup.title.string)
    return ""


def clean_text(text: str) -> str:
    """去除多餘空白但保留段落換行。"""
    text = text.replace("\xa0", " ").replace("　", " ").replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def node_to_text(node: Tag | BeautifulSoup) -> str:
    """把 HTML 節點轉成保留區塊換行的純文字。"""
    for tag in node.find_all(DROP_TAGS):
        tag.decompose()
    for br in node.find_all("br"):
        br.replace_with("\n")
    for tag in node.find_all(list(BLOCK_TAGS - {"br"})):
        tag.insert_before("\n")
        tag.insert_after("\n")
    return clean_text(node.get_text())


def html_to_text(html: str, selectors: list[str] | None = None) -> str:
    """依序嘗試 selectors，取第一個命中的節點；都沒有就取 <body>。"""
    soup = make_soup(html)
    target: Tag | BeautifulSoup | None = None
    for selector in selectors or []:
        found = soup.select_one(selector)
        if found is not None and clean_text(found.get_text()):
            target = found
            break
    if target is None:
        target = soup.body or soup
    return node_to_text(target)


def select_html(html: str, selectors: list[str]) -> str:
    """回傳第一個命中 selector 的節點 HTML（用來保存 raw_html 主內容區）。"""
    soup = make_soup(html)
    for selector in selectors:
        found = soup.select_one(selector)
        if found is not None:
            return str(found)
    return html


def normalize_label(label: str) -> str:
    """「學　　制」→「學制」、「申請(學)年度」→「申請(學)年度」。"""
    return re.sub(r"[\s　]+", "", label).strip("：:")


def extract_labeled_blocks(container: Tag, *, block_selector: str, label_selector: str, content_selector: str) -> dict[str, str]:
    """通用「標籤 + 內容」區塊抽取（圓夢助學網、海大生輔組都用這種版型）。"""
    result: dict[str, str] = {}
    for block in container.select(block_selector):
        label_node = block.select_one(label_selector)
        content_node = block.select_one(content_selector)
        if label_node is None or content_node is None:
            continue
        label = normalize_label(label_node.get_text())
        content = node_to_text(content_node)
        if label and content:
            result[label] = content
    return result


# ---- 附件取捨：哪些 PDF 值得下載下來併進原文（規範文件要，表單、預算書不要）
# 名稱像規範／說明文件：內容是資格、金額、申請方式的正式依據
DETAIL_ATTACHMENT_RE = re.compile(r"(要點|計畫|辦法|條例|規定|規範|簡章|須知|基準|標準|原則|注意事項|作業流程|申請說明|實施|核定本|公告|補助|補貼|津貼|給付|獎助|獎學金|助學|減免|救助|扶助|問答|Q&A|QA)", re.I)
# 名稱像表單／帳務文件：填寫用的空白表格或會計報表，沒有資格條件
FORM_ATTACHMENT_RE = re.compile(r"(申請書|申請表|申復表|報名表|切結書|同意書|委託書|承諾書|聲明書|申報表|領據|印領清冊|存摺|範本|範例|問卷|名冊|名單|一覽表|預算書|決算書|預算表|決算表|平衡表|現金流量|收支餘絀|餘絀撥補|統計表|對照表)")
# 名稱看不出內容（超連結文字只寫「pdf」「檔案下載」「請點此下載」）
MEANINGLESS_NAME_RE = re.compile(r"^[\s\[\(【（]*(?:pdf|odt|docx?|ods|xlsx?|csv|檔案?|附件|下載|點此|請點此下載參閱|標題|檔案下載|附件下載|pdf檔?下載?|\.\w+)[\s\]\)】）檔案下載參閱]*$", re.I)
# 本文短於這個長度時，連看不出名稱的附件也值得抓（資格與給付內容多半只寫在附件裡）
THIN_BODY_CHARS = 1500


def attachment_is_detail(name: str, *, body_chars: int = 0) -> bool:
    """這個附件值不值得下載：名稱像規範文件就抓；名稱看不出來時，只有在本文不足時才抓。表單與帳務文件一律不抓。"""
    name = (name or "").strip()
    if FORM_ATTACHMENT_RE.search(name):
        return False
    if DETAIL_ATTACHMENT_RE.search(name):
        return True
    return body_chars < THIN_BODY_CHARS


def meaningless_name(name: str) -> bool:
    """連結文字只有「pdf」「檔案下載」這類字眼或流水號檔名，不能當標題（中文檔名去掉副檔名後仍可用）。"""
    text = (name or "").strip()
    return not text or bool(MEANINGLESS_NAME_RE.match(text)) or bool(re.match(r"^[A-Za-z0-9_\-. ]+\.(?:pdf|docx?|odt|ods|xlsx?|csv)$", text, re.I))


def headline_title(text: str, *, max_length: int = 60) -> str:
    """PDF／附件的第一行有意義的文字當標題（連結文字是「pdf」時的替代來源）。"""
    for line in (text or "").splitlines():
        line = clean_text(line).strip(" :：-—　")
        if not line or meaningless_name(line):
            continue
        if re.match(r"^(發布單位|資料提供單位|承辦單位|聯絡電話|電話|傳真|附件|檔案大小)[：:]", line):
            continue
        if len(re.findall(r"[一-鿿]", line)) < 4:
            continue
        return line[:max_length]
    return ""


def find_attachments(container: Tag, base_url: str) -> list[dict]:
    attachments: list[dict] = []
    seen: set[str] = set()
    for anchor in container.find_all("a", href=True):
        href = anchor["href"].strip()
        if not ATTACHMENT_EXT.search(href):
            continue
        url = urljoin(base_url, href)
        if url in seen:
            continue
        seen.add(url)
        name = clean_text(anchor.get_text()) or url.rsplit("/", 1)[-1]
        ext = ATTACHMENT_EXT.search(href)
        attachments.append({"url": url, "name": name, "type": (ext.group(1).lower() if ext else "")})
    return attachments


def absolutize(base_url: str, href: str) -> str:
    return urljoin(base_url, href.replace("&amp;", "&"))


def pdf_to_text(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # pragma: no cover - 個別頁面解析失敗不影響整份
            continue
    return clean_text("\n".join(pages))


def docx_to_text(content: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return clean_text("\n".join(parts))
